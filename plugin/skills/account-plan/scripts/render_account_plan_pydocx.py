#!/usr/bin/env python3
"""
render_account_plan_pydocx.py — python-docx renderer for the account-plan skill.

Renders a declarative JSON block spec (see md_to_docx_spec.py, which converts this
skill's own markdown output into that spec) into a plain, CRO-readable working
document — no brand banner/colors, this never leaves the building. python-docx is
already a pyproject dependency; no npm/node needed.

Usage:
  uv run python scripts/md_to_docx_spec.py <plan.md> <spec.json>
  uv run python scripts/render_account_plan_pydocx.py <spec.json> <out.docx>

Reuses the four OOXML gotchas solved in the account-dossier skill's
render_dossier_pydocx.py (width units, strict child-element order, "start"/"end"
border sides, the zoom-percent template bug) — copied here rather than imported so
this skill stays self-contained (each skill's committed scripts are independently
reviewed; see CLAUDE.md "least-privilege by construction").

Adds inline markdown span parsing (**bold** and [text](url)) since an account
plan's prose and table cells use both heavily — the dossier renderer's blocks
never needed it because dossier specs are hand-composed with the bold/link
markers already routed to typed fields, but a markdown-derived spec keeps them
inline in plain text.
"""

import argparse
import json
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

# ── plain working-doc palette — no tenant brand, this document never leaves the building ──
INK = "1A1A1A"  # body text / headings
MUTED = "6B7280"  # meta text, footer
RULE = "334155"  # heading accent rule — dark slate, not a brand color
HEADER_FILL = "E5E7EB"  # table header row
LABEL_FILL = "F3F4F6"  # facts_table label column
BORDER = "CBD5E1"
FONT = "Calibri"
MONO_FONT = "Consolas"  # `code` spans — file paths appear throughout an account plan

MARGIN_DXA = 720
PAGE_W_DXA = 12240
PAGE_H_DXA = 15840
CONTENT_W_DXA = PAGE_W_DXA - 2 * MARGIN_DXA

# Table columns whose header matches one of these get extra relative width —
# narrative cells (an action, a gap, a mitigation) need room; a date/owner/score
# doesn't. Keeps a 6-7 column MEDDPICC/risk/action table from crushing its one
# long column into an unreadable sliver under equal-width division.
_WIDE_HEADER_HINTS = (
    "action",
    "detail",
    "where",
    "evidence",
    "win-result",
    "win result",
    "mitigation",
    "description",
    "summary",
    "gap-closing",
    "gap closing",
    "success criterion",
    "note",
    "read",
    "requirement",
    "exposure",
    "position",
    "counter",
    "criterion",
    "hook",
    "value",
)
_NARROW_HEADER_HINTS = (
    "date",
    "owner",
    "score",
    "priority",
    "status",
    "coverage",
    "rating",
    "likelihood",
    "impact",
    "#",
    "pts",
    "by",
)

_TCPR_ORDER = [
    "w:cnfStyle",
    "w:tcW",
    "w:gridSpan",
    "w:hMerge",
    "w:vMerge",
    "w:tcBorders",
    "w:shd",
    "w:noWrap",
    "w:tcMar",
    "w:textDirection",
    "w:tcFitText",
    "w:vAlign",
    "w:hideMark",
    "w:headers",
    "w:cellIns",
    "w:cellDel",
    "w:cellMerge",
    "w:tcPrChange",
]
_PPR_ORDER = [
    "w:pStyle",
    "w:keepNext",
    "w:keepLines",
    "w:pageBreakBefore",
    "w:framePr",
    "w:widowControl",
    "w:numPr",
    "w:suppressLineNumbers",
    "w:pBdr",
    "w:shd",
    "w:tabs",
    "w:suppressAutoHyphens",
    "w:kinsoku",
    "w:wordWrap",
    "w:overflowPunct",
    "w:topLinePunct",
    "w:autoSpaceDE",
    "w:autoSpaceDN",
    "w:bidi",
    "w:adjustRightInd",
    "w:snapToGrid",
    "w:spacing",
    "w:ind",
    "w:contextualSpacing",
    "w:mirrorIndents",
    "w:suppressOverlap",
    "w:jc",
    "w:textDirection",
    "w:textAlignment",
    "w:textboxTightWrap",
    "w:outlineLvl",
    "w:divId",
    "w:cnfStyle",
    "w:rPr",
    "w:sectPr",
    "w:pPrChange",
]
_TBLPR_ORDER = [
    "w:tblStyle",
    "w:tblpPr",
    "w:tblOverlap",
    "w:bidiVisual",
    "w:tblStyleRowBandSize",
    "w:tblStyleColBandSize",
    "w:tblW",
    "w:jc",
    "w:tblCellSpacing",
    "w:tblInd",
    "w:tblBorders",
    "w:shd",
    "w:tblLayout",
    "w:tblCellMar",
    "w:tblLook",
    "w:tblCaption",
    "w:tblDescription",
    "w:tblPrChange",
]
_BORDER_SIDES = ("top", "start", "bottom", "end")


def _insert_ordered(parent_elm, order, child_elm):
    tag = child_elm.tag
    local = tag.split("}", 1)[1] if "}" in tag else tag
    idx = order.index(f"w:{local}")
    successors = [
        f"w:{t.split(':', 1)[1]}" if not t.startswith("w:") else t for t in order[idx + 1 :]
    ]
    successor_qns = [qn(t) for t in successors]
    for existing in list(parent_elm):
        if existing.tag in successor_qns:
            existing.addprevious(child_elm)
            return child_elm
    parent_elm.append(child_elm)
    return child_elm


def hx(color):
    return str(color or "").lstrip("#").upper() or "000000"


def hp(size_halfpt):
    return Pt(size_halfpt / 2)


# ── inline markdown spans: **bold**, *italic*, `code`, [text](url) ─────────────
# Order matters: **bold** must be tried before *italic*, else the leading '*' of a
# bold marker matches as an italic delimiter and swallows the pair.
_SPAN_RE = re.compile(
    r"\*\*(?P<bold>.+?)\*\*"
    r"|(?<!\*)\*(?P<ital>[^*\n]+?)\*(?!\*)"
    r"|`(?P<code>[^`\n]+?)`"
    r"|\[(?P<ltext>[^\]]+)\]\((?P<lurl>[^)]+)\)"
)


def parse_inline(text):
    """Return a list of {text, bold, italic, mono, url} spans from a markdown string.

    Covers the inline markers this skill's own markdown actually emits. Anything
    unmatched passes through as literal text, so an unsupported marker degrades to
    visible punctuation rather than being dropped.
    """

    def span(t, bold=False, italic=False, mono=False, url=None):
        return {"text": t, "bold": bold, "italic": italic, "mono": mono, "url": url}

    spans = []
    pos = 0
    for m in _SPAN_RE.finditer(text or ""):
        if m.start() > pos:
            spans.append(span(text[pos : m.start()]))
        if m.group("bold") is not None:
            spans.append(span(m.group("bold"), bold=True))
        elif m.group("ital") is not None:
            spans.append(span(m.group("ital"), italic=True))
        elif m.group("code") is not None:
            spans.append(span(m.group("code"), mono=True))
        else:
            spans.append(span(m.group("ltext"), url=m.group("lurl")))
        pos = m.end()
    if pos < len(text or ""):
        spans.append(span(text[pos:]))
    return spans or [span(text or "")]


class Renderer:
    def __init__(self, spec, spec_dir):
        self.spec = spec
        self.spec_dir = spec_dir
        self.doc = Document()
        self._setup_document()

    def _setup_document(self):
        doc = self.doc
        normal = doc.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(11)
        normal.font.color.rgb = RGBColor.from_string(INK)

        section = doc.sections[0]
        section.page_width = Twips(PAGE_W_DXA)
        section.page_height = Twips(PAGE_H_DXA)
        section.top_margin = Twips(MARGIN_DXA)
        section.bottom_margin = Twips(MARGIN_DXA)
        section.left_margin = Twips(MARGIN_DXA)
        section.right_margin = Twips(MARGIN_DXA)

        self._fix_zoom_element()
        self._build_footer(section)

    def _fix_zoom_element(self):
        settings = self.doc.settings.element
        zoom = settings.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")

    def _build_footer(self, section):
        closing = self.spec.get("closingLine") or "Internal — not for customer distribution."
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{closing}   ·   Page ")
        r.font.name = FONT
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(MUTED)
        run = p.add_run()
        run.font.name = FONT
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_end)

    # ── low-level OOXML helpers (gotchas #1-3, see module docstring) ──────────
    def _set_cell_width(self, cell, dxa):
        tcPr = cell._tc.get_or_add_tcPr()
        tcW = tcPr.find(qn("w:tcW"))
        if tcW is None:
            tcW = OxmlElement("w:tcW")
            _insert_ordered(tcPr, _TCPR_ORDER, tcW)
        tcW.set(qn("w:w"), str(int(dxa)))
        tcW.set(qn("w:type"), "dxa")

    def _set_table_grid(self, table, column_widths_dxa):
        table.autofit = False
        tblPr = table._tbl.tblPr
        layout = tblPr.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            _insert_ordered(tblPr, _TBLPR_ORDER, layout)
        layout.set(qn("w:type"), "fixed")
        for i, col in enumerate(table.columns):
            col.width = Twips(column_widths_dxa[i])
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            _insert_ordered(tblPr, _TBLPR_ORDER, tblW)
        tblW.set(qn("w:w"), str(int(sum(column_widths_dxa))))
        tblW.set(qn("w:type"), "dxa")

    def _set_cell_shading(self, cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = tcPr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            _insert_ordered(tcPr, _TCPR_ORDER, shd)
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hx(fill_hex))

    def _set_cell_margins(self, cell, top=80, bottom=80, left=120, right=120):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = tcPr.find(qn("w:tcMar"))
        if tcMar is None:
            tcMar = OxmlElement("w:tcMar")
            _insert_ordered(tcPr, _TCPR_ORDER, tcMar)
        for side, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
            el = tcMar.find(qn(f"w:{side}"))
            if el is None:
                el = OxmlElement(f"w:{side}")
                tcMar.append(el)
            el.set(qn("w:w"), str(int(val)))
            el.set(qn("w:type"), "dxa")
        order = ["top", "left", "bottom", "right"]
        children = {c.tag.split("}", 1)[1]: c for c in list(tcMar)}
        for c in list(tcMar):
            tcMar.remove(c)
        for side in order:
            if side in children:
                tcMar.append(children[side])

    def _set_cell_borders(self, cell, color=BORDER, sz=4):
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = tcPr.find(qn("w:tcBorders"))
        if tcBorders is None:
            tcBorders = OxmlElement("w:tcBorders")
            _insert_ordered(tcPr, _TCPR_ORDER, tcBorders)
        for side in _BORDER_SIDES:
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), hx(color))
            existing = tcBorders.find(qn(f"w:{side}"))
            if existing is not None:
                tcBorders.remove(existing)
            tcBorders.append(el)

    def _add_heading_rule(self, paragraph, color, sz="10"):
        pPr = paragraph._p.get_or_add_pPr()
        pBdr = pPr.find(qn("w:pBdr"))
        if pBdr is None:
            pBdr = OxmlElement("w:pBdr")
            _insert_ordered(pPr, _PPR_ORDER, pBdr)
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), sz)
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), hx(color))
        pBdr.append(bottom)

    def _add_hyperlink_run(self, paragraph, url, text, size_halfpt, bold=False):
        part = paragraph.part
        r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), r_id)
        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        rStyle = OxmlElement("w:rStyle")
        rStyle.set(qn("w:val"), "Hyperlink")
        rPr.append(rStyle)
        rFonts = OxmlElement("w:rFonts")
        rFonts.set(qn("w:ascii"), FONT)
        rFonts.set(qn("w:hAnsi"), FONT)
        rPr.append(rFonts)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(size_halfpt))
        rPr.append(sz)
        if bold:
            rPr.append(OxmlElement("w:b"))
        r.append(rPr)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        hyperlink.append(r)
        paragraph._p.append(hyperlink)

    # ── text helpers ─────────────────────────────────────────────────────────
    def _run_plain(
        self, container, text, size=22, bold=False, italic=False, color=None, mono=False
    ):
        r = container.add_run("" if text is None else str(text))
        r.font.name = MONO_FONT if mono else FONT
        r.font.size = hp(size - 1 if mono else size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = RGBColor.from_string(hx(color) if color else INK)
        return r

    def _run_rich(self, paragraph, text, size=20, bold=False, italic=False, color=None):
        """Render `text` (may carry **bold**, *italic*, `code`, [link](url)) as one
        or more runs in an existing paragraph."""
        for span in parse_inline(text):
            span_bold = bold or span["bold"]
            span_italic = italic or span["italic"]
            if span["url"]:
                self._add_hyperlink_run(paragraph, span["url"], span["text"], size, bold=span_bold)
            else:
                self._run_plain(
                    paragraph,
                    span["text"],
                    size=size,
                    bold=span_bold,
                    italic=span_italic,
                    color=color,
                    mono=span["mono"],
                )

    def _paragraph(self, text, size=22, spacing_after=140, bold=False, italic=False, color=None):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Twips(spacing_after)
        self._run_rich(p, text, size=size, bold=bold, italic=italic, color=color)
        return p

    def _bullet(self, container, text, size=20, spacing_after=40):
        p = container.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Twips(spacing_after)
        self._run_rich(p, text, size=size)
        return p

    def _new_table(self, rows, cols, column_widths_dxa):
        table = self.doc.add_table(rows=rows, cols=cols)
        table.style = None
        self._set_table_grid(table, column_widths_dxa)
        return table

    def _column_widths(self, header, cols):
        """Weighted widths: headers matching a 'narrative' hint get 2x a plain
        column; headers matching a 'narrow' hint (date/owner/score/...) get 0.6x.
        Falls back to equal width with no header."""
        if not header:
            w = CONTENT_W_DXA // cols
            return [w] * cols
        weights = []
        for h in header:
            hl = (h or "").strip().lower()
            if any(hint in hl for hint in _WIDE_HEADER_HINTS):
                weights.append(2.0)
            elif any(hint in hl for hint in _NARROW_HEADER_HINTS):
                weights.append(0.6)
            else:
                weights.append(1.0)
        total = sum(weights)
        widths = [round(CONTENT_W_DXA * w / total) for w in weights]
        drift = CONTENT_W_DXA - sum(widths)
        widths[-1] += drift  # keep the table width exact
        return widths

    # ── block renderers ─────────────────────────────────────────────────────
    def render_title(self, blk):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Twips(60)
        self._run_plain(p, blk.get("text", ""), size=48, bold=True, color=INK)
        if blk.get("subtitle"):
            p2 = self.doc.add_paragraph()
            p2.paragraph_format.space_after = Twips(40)
            self._run_rich(p2, blk["subtitle"], size=20, color=MUTED)
        if blk.get("meta"):
            p3 = self.doc.add_paragraph()
            p3.paragraph_format.space_after = Twips(20)
            self._add_heading_rule(p3, RULE, sz="16")
            self._run_rich(p3, blk["meta"], size=18, italic=True, color=MUTED)
        else:
            rule_p = self.doc.add_paragraph()
            rule_p.paragraph_format.space_after = Twips(20)
            self._add_heading_rule(rule_p, RULE, sz="16")
        self.doc.add_paragraph().paragraph_format.space_after = Twips(80)

    def render_heading(self, blk):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Twips(280)
        p.paragraph_format.space_after = Twips(120)
        self._add_heading_rule(p, RULE)
        self._run_plain(p, blk.get("text", ""), size=27, bold=True, color=INK)

    def render_subheading(self, blk):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Twips(160)
        p.paragraph_format.space_after = Twips(80)
        self._run_rich(p, blk.get("text", ""), size=22, bold=True, color=INK)

    def render_paragraph(self, blk):
        self._paragraph(
            blk.get("text", ""), size=21, spacing_after=140, italic=bool(blk.get("italic"))
        )

    def render_bullets(self, blk):
        for item in blk.get("items") or []:
            self._bullet(self.doc, item, size=20)
        self.doc.add_paragraph().paragraph_format.space_after = Twips(60)

    def render_questions(self, blk):
        for g in blk.get("groups") or []:
            p = self.doc.add_paragraph()
            p.paragraph_format.space_before = Twips(100)
            p.paragraph_format.space_after = Twips(60)
            self._run_plain(p, g.get("title", ""), size=21, bold=True, color=INK)
            for q in g.get("items") or []:
                self._bullet(self.doc, q, size=20)

    def render_facts_table(self, blk):
        rows = blk.get("rows") or []
        if not rows:
            return
        w0 = round(CONTENT_W_DXA * 0.28)
        w1 = CONTENT_W_DXA - w0
        table = self._new_table(len(rows), 2, [w0, w1])
        for r, pair in enumerate(rows):
            label_cell, value_cell = table.rows[r].cells
            self._set_cell_shading(label_cell, LABEL_FILL)
            for cell, dxa in ((label_cell, w0), (value_cell, w1)):
                self._set_cell_margins(cell)
                self._set_cell_borders(cell)
                self._set_cell_width(cell, dxa)
            self._run_rich(
                label_cell.paragraphs[0], pair[0] if len(pair) > 0 else "", size=20, bold=True
            )
            self._run_rich(value_cell.paragraphs[0], pair[1] if len(pair) > 1 else "", size=20)
        self.doc.add_paragraph().paragraph_format.space_after = Twips(120)

    def render_table(self, blk):
        header = blk.get("header") or []
        body = blk.get("rows") or []
        cols = len(header) or (len(body[0]) if body else 1)
        if cols == 0:
            return
        widths = self._column_widths(header, cols)
        total_rows = (1 if header else 0) + len(body)
        if total_rows == 0:
            return
        size = 18 if cols >= 6 else 20
        table = self._new_table(total_rows, cols, widths)
        r_idx = 0
        if header:
            for c, text in enumerate(header):
                cell = table.rows[0].cells[c]
                self._set_cell_shading(cell, HEADER_FILL)
                self._set_cell_margins(cell)
                self._set_cell_borders(cell)
                self._set_cell_width(cell, widths[c])
                self._run_plain(cell.paragraphs[0], text, size=size, bold=True, color=INK)
            r_idx = 1
        for row in body:
            for c in range(cols):
                text = row[c] if c < len(row) else ""
                cell = table.rows[r_idx].cells[c]
                self._set_cell_margins(cell)
                self._set_cell_borders(cell)
                self._set_cell_width(cell, widths[c])
                self._run_rich(cell.paragraphs[0], text, size=size)
            r_idx += 1
        self.doc.add_paragraph().paragraph_format.space_after = Twips(120)

    def render_spacer(self, _blk):
        self.doc.add_paragraph().paragraph_format.space_after = Twips(80)

    def render_pagebreak(self, _blk):
        self.doc.add_page_break()

    # ── entry point ──────────────────────────────────────────────────────────
    def render(self):
        dispatch = {
            "title": self.render_title,
            "heading": self.render_heading,
            "subheading": self.render_subheading,
            "paragraph": self.render_paragraph,
            "bullets": self.render_bullets,
            "questions": self.render_questions,
            "facts_table": self.render_facts_table,
            "table": self.render_table,
            "spacer": self.render_spacer,
            "pagebreak": self.render_pagebreak,
        }
        blocks = self.spec.get("blocks") or []
        count = 0
        for blk in blocks:
            if not isinstance(blk, dict):
                continue
            fn = dispatch.get(blk.get("type"))
            if fn is None:
                print(f"unknown block type, skipping: {blk.get('type')}", file=sys.stderr)
                continue
            fn(blk)
            count += 1
        if not blocks:
            self._paragraph("(empty spec)")
        return count


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("spec_path")
    parser.add_argument("out_path")
    args = parser.parse_args()

    with open(args.spec_path, encoding="utf-8") as f:
        spec = json.load(f)

    renderer = Renderer(spec, os.path.dirname(os.path.abspath(args.spec_path)))
    count = renderer.render()
    renderer.doc.save(args.out_path)
    size = os.path.getsize(args.out_path)
    print(f"wrote {args.out_path} ({size} bytes, {count} blocks)")


if __name__ == "__main__":
    main()
